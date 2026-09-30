"""
Script de generación del Pitch Deck Comercial y Técnico para ENIAC Labs.
Diseñado con las directrices de 'ui-ux-pro-max' y 'pptx-builder':
- Formato 16:9 Widescreen (13.333" x 7.5")
- Paleta Light B2B Tech: Navy profundo (#0F172A), Teal (#0D9488), Slate suave (#F8FAFC)
- Slide 1: Portada con layout split 60/40 (Info institucional + Team / Hero image 3D)
- Slides 2-6: Contenido estructurado en cards modulares con badges y cero muros de texto
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Layout en blanco

    # Paleta cromática UI/UX Pro Max (Enterprise B2B Cloud)
    COLOR_BG = RGBColor(248, 250, 252)        # #F8FAFC
    COLOR_PRIMARY = RGBColor(15, 23, 42)       # #0F172A (Navy)
    COLOR_TEAL = RGBColor(13, 148, 136)        # #0D9488 (Teal)
    COLOR_BLUE_ACCENT = RGBColor(2, 132, 199)  # #0284C7 (Sky Blue)
    COLOR_PURPLE = RGBColor(99, 102, 241)      # #6366F1 (Indigo)
    COLOR_CARD = RGBColor(255, 255, 255)       # #FFFFFF
    COLOR_BORDER = RGBColor(226, 232, 240)     # #E2E8F0
    COLOR_TEXT = RGBColor(30, 41, 59)          # #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B
    COLOR_BADGE_BG = RGBColor(240, 253, 250)   # #F0FDFA
    COLOR_EMERALD = RGBColor(16, 185, 129)     # #10B981

    FONT_FAMILY = "Segoe UI"

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, category_badge, title, subtitle):
        # Category Badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(3.2), Inches(0.38)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_BADGE_BG
        badge.line.color.rgb = RGBColor(204, 251, 241)
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_badge.upper()
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL
        p.alignment = PP_ALIGN.CENTER

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.88), Inches(11.7), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_t = tf_title.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_s = tf_sub.paragraphs[0]
        p_s.text = subtitle
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 1: PORTADA (Split 60 / 40)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # 1. Encabezado institucional discreto arriba
    inst_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_inst = inst_box.text_frame
    p_inst = tf_inst.paragraphs[0]
    p_inst.text = "UNIVERSIDAD PERUANA UNIÓN  •  INGENIERÍA DE SISTEMAS  •  ARQUITECTURA DE MICROSERVICIOS 2026"
    p_inst.font.name = FONT_FAMILY
    p_inst.font.size = Pt(9.5)
    p_inst.font.bold = True
    p_inst.font.color.rgb = COLOR_MUTED

    # Línea decorativa sutil
    sep = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.02))
    sep.fill.solid()
    sep.fill.fore_color.rgb = COLOR_BORDER
    sep.line.fill.background()

    # Columna Izquierda (60% ~ 7.2 pulgadas)
    # Logo Pill Badge
    logo_badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(2.3), Inches(0.42))
    logo_badge.fill.solid()
    logo_badge.fill.fore_color.rgb = COLOR_PRIMARY
    logo_badge.line.fill.background()
    p_l = logo_badge.text_frame.paragraphs[0]
    p_l.text = "⚡ ENIAC LABS"
    p_l.font.name = FONT_FAMILY
    p_l.font.bold = True
    p_l.font.size = Pt(12)
    p_l.font.color.rgb = RGBColor(255, 255, 255)
    p_l.alignment = PP_ALIGN.CENTER

    # Título Principal
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(6.8), Inches(1.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "Plataforma Distribuida de Gestión y Venta de Hardware"
    p_t1.font.name = FONT_FAMILY
    p_t1.font.size = Pt(28)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_PRIMARY

    # Tagline corto
    tag_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(6.8), Inches(0.45))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "“Aprender, innovar y construir infraestructura que escala”"
    p_tag.font.name = FONT_FAMILY
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.italic = True
    p_tag.font.color.rgb = COLOR_TEAL

    # Breve descripción comercial
    desc_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(6.8), Inches(1.2))
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = (
        "Ecosistema cloud-native de alto rendimiento con arquitectura hexagonal, orquestación "
        "resiliente Saga y seguridad perimetral Zero-Trust para cotización, control atómico de inventario "
        "y procesamiento de pagos en tiempo real."
    )
    p_desc.font.name = FONT_FAMILY
    p_desc.font.size = Pt(11.5)
    p_desc.font.color.rgb = COLOR_TEXT

    # Bloque de Equipo y Docente (Card suave)
    team_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(6.8), Inches(1.75))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = COLOR_CARD
    team_card.line.color.rgb = COLOR_BORDER
    team_card.line.width = Pt(1)

    tf_tc = team_card.text_frame
    tf_tc.word_wrap = True
    tf_tc.margin_left = Inches(0.25)
    tf_tc.margin_top = Inches(0.18)

    p_th = tf_tc.paragraphs[0]
    p_th.text = "EQUIPO DE INGENIERÍA & GOBERNANZA"
    p_th.font.name = FONT_FAMILY
    p_th.font.bold = True
    p_th.font.size = Pt(9.5)
    p_th.font.color.rgb = COLOR_TEAL

    p_t1 = tf_tc.add_paragraph()
    p_t1.text = "• Desarrolladores Principales: Equipo ENIAC Labs (Ingeniería de Sistemas)"
    p_t1.font.name = FONT_FAMILY
    p_t1.font.size = Pt(10.5)
    p_t1.font.color.rgb = COLOR_PRIMARY

    p_t2 = tf_tc.add_paragraph()
    p_t2.text = "• Docente Asesor: Mag. Especialista en Sistemas Distribuidos"
    p_t2.font.name = FONT_FAMILY
    p_t2.font.size = Pt(10.5)
    p_t2.font.color.rgb = COLOR_PRIMARY

    p_t3 = tf_tc.add_paragraph()
    p_t3.text = "• Stack: Spring Boot 3.3/Cloud • PostgreSQL • Docker • JWT Stateless • Resilience4j"
    p_t3.font.name = FONT_FAMILY
    p_t3.font.size = Pt(9.5)
    p_t3.font.color.rgb = COLOR_MUTED

    # Columna Derecha (40% ~ 4.7 pulgadas): Hero Image
    hero_path = r"c:\Users\USUARIO\Documents\micro\presentations\assets\eniaclabs_hero.jpg"
    if os.path.exists(hero_path):
        # Card contenedor para imagen
        frame_shape = slide1.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.1), Inches(4.5), Inches(5.75)
        )
        frame_shape.fill.solid()
        frame_shape.fill.fore_color.rgb = RGBColor(241, 245, 249)
        frame_shape.line.color.rgb = COLOR_BORDER
        frame_shape.line.width = Pt(1.5)

        # Imagen insertada
        slide1.shapes.add_picture(hero_path, Inches(8.1), Inches(1.2), Inches(4.3), Inches(5.55))

    # ==========================================
    # SLIDE 2: ARQUITECTURA DISTRIBUIDA
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(
        slide2,
        "Diseño de Sistemas • Topología Cloud-Native",
        "Arquitectura de Microservicios & Desacoplamiento",
        "Estructura orientada a servicios desacoplados con base de datos independiente por dominio."
    )

    cards_data_s2 = [
        {
            "icon": "🌐",
            "badge": "EDGE GATEWAY",
            "title": "Spring Cloud Gateway",
            "points": [
                "Punto único de entrada perimetral con enrutamiento reactivo dinámico.",
                "Filtros de validación de identidad y propagación de credenciales hacia downstream.",
                "Balanceo de carga del lado del cliente integrado con Netflix Eureka."
            ],
            "accent": COLOR_TEAL
        },
        {
            "icon": "🗄️",
            "badge": "DATABASE PER SERVICE",
            "title": "Aislamiento de Persistencia",
            "points": [
                "Bases de datos PostgreSQL dedicadas para Catálogo, Órdenes, Pagos y Auth.",
                "Sin dependencias cruzadas en SQL ni acoplamiento de esquema relacional.",
                "Migraciones deterministas y versionadas independientes para cada dominio."
            ],
            "accent": COLOR_BLUE_ACCENT
        },
        {
            "icon": "⚙️",
            "badge": "CENTRALIZED OPS",
            "title": "Eureka & Config Server",
            "points": [
                "Auto-registro y descubrimiento transparente de instancias en runtime.",
                "Centralización de perfiles 'dev' y 'prod' en repositorio Git unificado.",
                "Desacoplamiento total de I/O de red fuera de transacciones de base de datos."
            ],
            "accent": COLOR_PURPLE
        }
    ]

    card_width = Inches(3.68)
    card_height = Inches(4.7)
    card_y = Inches(2.1)

    for i, c in enumerate(cards_data_s2):
        card_x = Inches(0.8 + i * (3.68 + 0.34))
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        # Header badge inside card
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.25)

        p0 = tf_c.paragraphs[0]
        p0.text = f"{c['icon']}  {c['badge']}"
        p0.font.name = FONT_FAMILY
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = c['accent']

        p_title = tf_c.add_paragraph()
        p_title.text = c['title']
        p_title.font.name = FONT_FAMILY
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(14)

        for pt in c['points']:
            p_pt = tf_c.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(11)
            p_pt.font.color.rgb = COLOR_TEXT
            p_pt.space_after = Pt(10)

    # ==========================================
    # SLIDE 3: RESILIENCIA & SAGA PATTERN
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(
        slide3,
        "Transaccionalidad Distribuida • Alta Concurrencia",
        "Resiliencia & Transacciones Atómicas (Saga Pattern)",
        "Garantía de consistencia eventual ante caídas de red y eliminación de race conditions."
    )

    cards_data_s3 = [
        {
            "icon": "🛡️",
            "badge": "CIRCUIT BREAKER",
            "title": "Tolerancia a Fallos",
            "points": [
                "Circuit breakers con Resilience4j en comunicación síncrona Feign/REST.",
                "Timeouts configurados: 2000ms respuesta y 3500ms ventana global.",
                "Degradación elegante que previene cascadas de fallos en el cluster."
            ],
            "accent": COLOR_TEAL
        },
        {
            "icon": "⚡",
            "badge": "ATOMIC CONCURRENCY",
            "title": "Control Concurrente",
            "points": [
                "Actualizaciones atómicas 'UPDATE WHERE stock >= :cant' a nivel de fila.",
                "Prevención matemática de sobreventa (overselling) bajo alto tráfico.",
                "Cero bloqueos pesados ni problemas de lectura sucia en Catálogo."
            ],
            "accent": COLOR_EMERALD
        },
        {
            "icon": "🔄",
            "badge": "SAGA COMPENSATION",
            "title": "Compensación Automática",
            "points": [
                "Patrón Saga Coreografiado/Orquestado para rollback transaccional.",
                "Restitución inmediata de stock si el pago es rechazado o cancelado.",
                "Reconciliación idempotente para eventos asíncronos y webhooks."
            ],
            "accent": COLOR_PURPLE
        }
    ]

    for i, c in enumerate(cards_data_s3):
        card_x = Inches(0.8 + i * (3.68 + 0.34))
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.25)

        p0 = tf_c.paragraphs[0]
        p0.text = f"{c['icon']}  {c['badge']}"
        p0.font.name = FONT_FAMILY
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = c['accent']

        p_title = tf_c.add_paragraph()
        p_title.text = c['title']
        p_title.font.name = FONT_FAMILY
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(14)

        for pt in c['points']:
            p_pt = tf_c.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(11)
            p_pt.font.color.rgb = COLOR_TEXT
            p_pt.space_after = Pt(10)

    # ==========================================
    # SLIDE 4: SEGURIDAD PERIMETRAL JWT
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(
        slide4,
        "Seguridad Distribuida • Zero-Trust Perimeter",
        "Seguridad Perimetral & Propagación JWT",
        "Validación criptográfica centralizada en Gateway con confianza interna verificada."
    )

    cards_data_s4 = [
        {
            "icon": "🔐",
            "badge": "RESOURCE SERVER",
            "title": "Gateway como Escudo",
            "points": [
                "Spring Security stateless con validación estricta de firmas HMAC/RSA.",
                "Retorno inmediato de HTTP 401 para peticiones no autenticadas.",
                "Bloqueo HTTP 403 para usuarios sin privilegios según rol (ADMIN/CLIENTE)."
            ],
            "accent": COLOR_TEAL
        },
        {
            "icon": "🏷️",
            "badge": "CLAIM INJECTION",
            "title": "Propagación de Claims",
            "points": [
                "Extracción segura de claims del token en el Edge Gateway.",
                "Inyección en cabeceras confiables downstream: X-User-Id, X-User-Roles.",
                "Eliminación de campos mutables en DTOs para evitar suplantación de identidad."
            ],
            "accent": COLOR_BLUE_ACCENT
        },
        {
            "icon": "🛡️",
            "badge": "DEFENSE-IN-DEPTH",
            "title": "Aislamiento de Red",
            "points": [
                "Microservicios internos no expuestos directamente a redes públicas.",
                "Auditoría y filtros CORS granulares con prevención de ataques CSRF.",
                "Gestión de ciclo de vida de tokens con rotación y expiración estricta."
            ],
            "accent": COLOR_EMERALD
        }
    ]

    for i, c in enumerate(cards_data_s4):
        card_x = Inches(0.8 + i * (3.68 + 0.34))
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.25)

        p0 = tf_c.paragraphs[0]
        p0.text = f"{c['icon']}  {c['badge']}"
        p0.font.name = FONT_FAMILY
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = c['accent']

        p_title = tf_c.add_paragraph()
        p_title.text = c['title']
        p_title.font.name = FONT_FAMILY
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(14)

        for pt in c['points']:
            p_pt = tf_c.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(11)
            p_pt.font.color.rgb = COLOR_TEXT
            p_pt.space_after = Pt(10)

    # ==========================================
    # SLIDE 5: OBSERVABILIDAD & TRAZABILIDAD
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(
        slide5,
        "Telemetría & Monitoreo • SRE Best Practices",
        "Trazabilidad Distribuida & Observabilidad",
        "Visibilidad extremo a extremo de cada petición a través de la malla de microservicios."
    )

    cards_data_s5 = [
        {
            "icon": "📍",
            "badge": "CORRELATION TRACING",
            "title": "X-Correlation-Id End-to-End",
            "points": [
                "Generación automática en Gateway o captura de cabecera de entrada.",
                "Inyección permanente en MDC (Mapped Diagnostic Context) de Logback.",
                "Propagación obligatoria en cada salto HTTP vía interceptores Feign/RestTemplate."
            ],
            "accent": COLOR_TEAL
        },
        {
            "icon": "🩺",
            "badge": "HEALTH PROBES",
            "title": "Actuator & K8s Probes",
            "points": [
                "Endpoints de salud estandarizados en /actuator/health para cada nodo.",
                "Separación de sondas Liveness (supervivencia) y Readiness (tráfico).",
                "Monitoreo proactivo del pool de conexiones HikariCP y PostgreSQL."
            ],
            "accent": COLOR_BLUE_ACCENT
        },
        {
            "icon": "📊",
            "badge": "AUDIT & LOGS",
            "title": "Auditoría Centralizada",
            "points": [
                "Logs estructurados en JSON con timestamp, correlationId y microservicio.",
                "Diagnóstico acelerado: rastreo de cualquier fallo en menos de 30 segundos.",
                "Base lista para ingestión en stacks OpenTelemetry / Loki / Datadog."
            ],
            "accent": COLOR_PURPLE
        }
    ]

    for i, c in enumerate(cards_data_s5):
        card_x = Inches(0.8 + i * (3.68 + 0.34))
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.25)

        p0 = tf_c.paragraphs[0]
        p0.text = f"{c['icon']}  {c['badge']}"
        p0.font.name = FONT_FAMILY
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = c['accent']

        p_title = tf_c.add_paragraph()
        p_title.text = c['title']
        p_title.font.name = FONT_FAMILY
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(14)

        for pt in c['points']:
            p_pt = tf_c.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(11)
            p_pt.font.color.rgb = COLOR_TEXT
            p_pt.space_after = Pt(10)

    # ==========================================
    # SLIDE 6: MÉTRICAS DE IMPACTO & VALOR COMERCIAL
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(
        slide6,
        "Resultados Comerciales • Valor Tecnológico",
        "Impacto del Negocio & Eficiencia Operativa",
        "Beneficios cuantificables obtenidos tras la modernización y optimización de arquitectura."
    )

    stats = [
        {"val": "99.9%", "label": "Disponibilidad SLA", "desc": "Tolerancia a fallos mediante Circuit Breakers y pools aislados."},
        {"val": "-75%", "label": "Latencia de Checkout", "desc": "Desacoplamiento de I/O de red de las transacciones de BD."},
        {"val": "0", "label": "Riesgo de Sobreventa", "desc": "Consistencia estricta gracias a queries atómicas condicionales."},
        {"val": "100%", "label": "Trazabilidad de Flujo", "desc": "Visibilidad completa con X-Correlation-Id en cada hop del sistema."}
    ]

    stat_width = Inches(2.7)
    stat_height = Inches(4.7)
    for i, s in enumerate(stats):
        stat_x = Inches(0.8 + i * (2.7 + 0.3))
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, stat_x, card_y, stat_width, stat_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf_s = card.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.2)
        tf_s.margin_right = Inches(0.2)
        tf_s.margin_top = Inches(0.5)

        p_num = tf_s.paragraphs[0]
        p_num.text = s["val"]
        p_num.font.name = FONT_FAMILY
        p_num.font.bold = True
        p_num.font.size = Pt(36)
        p_num.font.color.rgb = COLOR_TEAL
        p_num.alignment = PP_ALIGN.CENTER
        p_num.space_after = Pt(10)

        p_lbl = tf_s.add_paragraph()
        p_lbl.text = s["label"]
        p_lbl.font.name = FONT_FAMILY
        p_lbl.font.bold = True
        p_lbl.font.size = Pt(14)
        p_lbl.font.color.rgb = COLOR_PRIMARY
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.space_after = Pt(14)

        p_d = tf_s.add_paragraph()
        p_d.text = s["desc"]
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_MUTED
        p_d.alignment = PP_ALIGN.CENTER

    output_path = r"c:\Users\USUARIO\Documents\micro\presentations\ENIAC_Labs_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
