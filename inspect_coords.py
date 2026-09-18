from pptx import Presentation

prs = Presentation(r"C:\Users\USUARIO\Documents\KIDO.pptx")
for idx, slide in enumerate(prs.slides):
    print(f"=== Slide {idx+1} ({len(slide.shapes)} shapes) ===")
    for s_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace("\n", " ")[:50]
        elif shape.has_table:
            txt = f"Table {len(shape.table.rows)}x{len(shape.table.columns)}"
        elif shape.shape_type == 13:
            txt = f"Picture ({shape.image.content_type})"
        print(f"  [{s_idx}] {shape.name} | Type={shape.shape_type} | Pos=({shape.left},{shape.top},{shape.width},{shape.height}) | {txt}")
