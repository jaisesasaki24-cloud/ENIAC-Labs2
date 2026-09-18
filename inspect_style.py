from pptx import Presentation

prs = Presentation(r"C:\Users\USUARIO\Documents\KIDO.pptx")
slide = prs.slides[0]
print("Slide 1 background:", slide.background.fill.type if slide.background else "None")
for s in slide.shapes:
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                font_color = r.font.color.rgb if r.font.color and hasattr(r.font.color, "rgb") else "default"
                print(f"Text: '{r.text[:30]}' Font: {r.font.name} Size: {r.font.size} Color: {font_color}")
