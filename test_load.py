import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

source_pptx = r"C:\Users\USUARIO\Documents\KIDO.pptx"
output_pptx = r"C:\Users\USUARIO\Documents\ENIAC_Labs_Presentacion.pptx"

prs = Presentation(source_pptx)
print("Loaded KIDO.pptx with", len(prs.slides), "slides")
