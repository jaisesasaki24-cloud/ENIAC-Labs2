from pptx import Presentation

prs = Presentation(r'C:\Users\USUARIO\Documents\KIDO.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            lines = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if lines:
                print(f"  [{s_idx}] {shape.name}: { ' | '.join(lines) }")
        elif shape.has_table:
            print(f"  [{s_idx}] TABLE {shape.name}:")
            for r in shape.table.rows:
                row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print("      | " + " | ".join(row_vals) + " |")
