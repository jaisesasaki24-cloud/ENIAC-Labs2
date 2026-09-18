from pptx import Presentation

prs = Presentation(r"C:\Users\USUARIO\Documents\KIDO.pptx")
for i in [1, 2, 3, 6, 9, 16]:
    slide = prs.slides[i]
    print(f"\n--- Slide {i+1} ---")
    for s in slide.shapes:
        fill_type = s.fill.type if hasattr(s, "fill") and s.fill else "None"
        line_color = s.line.color.rgb if hasattr(s, "line") and s.line and s.line.color and hasattr(s.line.color, "rgb") else "none"
        text = s.text[:40].replace("\n", " ") if s.has_text_frame else ""
        print(f"Shape: {s.name} | Type: {s.shape_type} | Fill: {fill_type} | Line: {line_color} | Text: '{text}'")
